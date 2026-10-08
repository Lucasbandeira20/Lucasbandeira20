"""Passo 1: separa os SKUs novos da exportação do TikTok.

Uso: python3 -I novos.py <exportacao.xlsx> <base_G.json> <controle_C.json> <pasta_lote>

- exportacao.xlsx: arquivo all_information_template do TikTok Seller Center (aba Template).
- base_G.json: get_values de 'BASE PRODUTOS TIKTOK'!G1:G (IDs de SKU já na BASE).
- controle_C.json: get_values de 'CONTROLE ANÚNCIOS TIKTOKSHOP 03.10'!C1:C (SKUs Bling no CONTROLE).

Gera na pasta do lote:
- novos.json: linhas novas com as 20 colunas da BASE (A..T), na ordem da exportação.
- skus_todos.json: SKUs Bling (seller_sku) de todas as linhas novas.
- skus_criar.json: SKUs Bling que não estão no CONTROLE (precisam de custo e linha nova).
"""
import collections
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import load_values  # noqa: E402

exp, base_g, ctl_c, out = sys.argv[1:5]
os.makedirs(out, exist_ok=True)

KEYS = ['product_id', 'category', 'product_name', 'product_status', 'gtin_type', 'gtin_code', 'sku_id',
        'variation_value', 'brand', 'price', 'seller_sku', 'minimum_order_quantity', 'maximum_order_quantity',
        'cumulative_order_quantity', 'parcel_weight', 'parcel_length', 'parcel_width', 'parcel_height', 'cod',
        'size_chart']

wb = openpyxl.load_workbook(exp, data_only=True)
ws = wb['Template'] if 'Template' in wb.sheetnames else wb.active
rows = [[c.value for c in r] for r in ws.iter_rows()]
H = rows[0]
ix = {h: i for i, h in enumerate(H) if h}
falt = [k for k in KEYS if k not in ix]
if falt:
    print('ATENÇÃO, colunas não encontradas na exportação:', falt)
dados = [[r[ix[k]] if k in ix else '' for k in KEYS] for r in rows[5:] if r[0] not in (None, '')]


def conv(r):
    o = []
    for j, x in enumerate(r):
        if x is None:
            x = ''
        if j in (0, 5, 6, 10):
            x = str(x).strip()
            if x.endswith('.0'):
                x = x[:-2]
        elif j == 9:
            x = float(x) if x != '' else ''
        elif j in (14, 15, 16, 17) and x != '':
            x = int(float(x))
        o.append(x)
    return o


dados = [conv(r) for r in dados]
na_base = {str(r[0]).strip() for r in load_values(base_g)[1:] if r and r[0]}
no_ctl = {str(r[0]).strip().lstrip('0') for r in load_values(ctl_c)[1:] if r and r[0]}

novos = [r for r in dados if r[6] not in na_base]
sumiram = len(na_base - {r[6] for r in dados})
criar = list(dict.fromkeys(r[10] for r in novos if r[10].lstrip('0') not in no_ctl))
todos = list(dict.fromkeys(r[10] for r in novos))

json.dump(novos, open(f'{out}/novos.json', 'w'), ensure_ascii=False)
json.dump(todos, open(f'{out}/skus_todos.json', 'w'))
json.dump(criar, open(f'{out}/skus_criar.json', 'w'))

print(f'exportação: {len(dados)} SKUs | já na BASE: {len(dados) - len(novos)} | NOVOS: {len(novos)} '
      f'({len({r[0] for r in novos})} produtos) | sumiram da exportação: {sumiram}')
print(f'novos já no CONTROLE: {sum(1 for r in novos if r[10] not in criar)} | a criar no CONTROLE: {len(criar)} SKUs únicos')
print('status:', dict(collections.Counter(r[3] for r in novos)))
print('sem seller_sku:', [r[2] for r in novos if not r[10]])
dup = [k for k, v in collections.Counter(r[10] for r in dados).items() if v > 1 and k]
if dup:
    print('SKU Bling em mais de um anúncio:', dup)
for r in novos:
    print(('CRIAR ' if r[10] in criar else 'CTL   ') + r[10], '|', r[2][:60], '|', r[7], '|', r[9])
