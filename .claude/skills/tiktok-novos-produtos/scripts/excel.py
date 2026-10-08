"""Passo 5: gera o Excel de promoção (modelo FixedPriceWithSKU do TikTok), só colunas A a E.

Uso: python3 -I excel.py <promocoes_A_C.json> <saida.xlsx>

promocoes_A_C.json: get_values de 'PROMOÇÕES TIKTOK'!A<ini>:C<fim> do lote (já com valores colados).
Linhas sem preço (SKUs fora da promoção) são ignoradas.
"""
import os
import shutil
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import load_values  # noqa: E402

src, out = sys.argv[1], sys.argv[2]
tpl = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'assets',
                   'modelo_promocao_FixedPriceWithSKU.xlsx')
v = [r for r in load_values(src) if len(r) >= 3 and r[2]]
erros = [r for r in v if len(str(r[0])) != 19 or len(str(r[1])) != 19 or ',' in str(r[2])]
if erros:
    sys.exit(f'Linhas com ID fora de 19 dígitos ou preço com vírgula: {erros[:5]}')

shutil.copy(tpl, out)
wb = openpyxl.load_workbook(out)
ws = wb.active
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for c in row:
        c.value = None
for row in ws.iter_rows(min_row=1, max_row=1):
    for c in row[5:]:
        c.value = None
for i, r in enumerate(v, start=2):
    for j, val in enumerate(r[:3], start=1):
        c = ws.cell(i, j)
        c.value = str(val)
        c.number_format = '@'
wb.save(out)
print(f'{out}: {len(v)} SKUs')
